from __future__ import annotations

import random
from pathlib import Path

from docx import Document
from docx.shared import RGBColor


KEYWORDS = {
    "实验目的": 10,
    "实验原理": 10,
    "实验内容": 10,
    "程序": 10,
    "结果": 10,
    "总结": 10,
    "心得": 10,
}


def evaluate_report(docx_path: Path) -> tuple[float, str]:
    """Simulate a large model evaluation.

    The evaluation searches for pre-defined keywords and awards points based on their presence.
    A small random adjustment is added to avoid identical scores.
    """

    document = Document(docx_path)
    full_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    score = 60
    matched_sections: list[str] = []
    for keyword, weight in KEYWORDS.items():
        if keyword in full_text:
            score += weight
            matched_sections.append(keyword)
    score = min(score + random.uniform(-5, 5), 100)
    score = round(score, 1)

    if matched_sections:
        feedback = f"报告结构完整，涵盖：{', '.join(matched_sections)}，请继续保持。"
    else:
        feedback = "建议补充实验目的、过程、结果等内容，让报告更加完整。"
    return score, feedback


def annotate_report(docx_path: Path, score: float, feedback: str, output_path: Path) -> Path:
    document = Document(docx_path)
    if document.paragraphs:
        target_paragraph = document.paragraphs[0].insert_paragraph_before()
    else:
        target_paragraph = document.add_paragraph()

    score_run = target_paragraph.add_run(f"成绩：{score}分")
    score_run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
    score_run.font.bold = True

    target_paragraph.add_run().add_break()

    feedback_run = target_paragraph.add_run(f"评语：{feedback}")
    feedback_run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
    feedback_run.font.bold = True

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path
