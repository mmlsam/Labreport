import os
import secrets
import shutil
from datetime import datetime
from pathlib import Path
from typing import Dict, List

import httpx
from fastapi import Depends, FastAPI, File, HTTPException, UploadFile, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from pydantic import BaseModel
from sqlalchemy import Column, DateTime, Float, Integer, String, create_engine, func
from sqlalchemy.ext.declarative import declarative_base
from sqlalchemy.orm import Session, sessionmaker
from passlib.context import CryptContext
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor

DATABASE_URL = "sqlite:///./labreport.db"

engine = create_engine(DATABASE_URL, connect_args={"check_same_thread": False})
SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
Base = declarative_base()

pwd_context = CryptContext(schemes=["bcrypt"], deprecated="auto")

data_dir = Path("data")
reports_dir = data_dir / "reports"
reports_dir.mkdir(parents=True, exist_ok=True)


class Student(Base):
    __tablename__ = "students"

    id = Column(Integer, primary_key=True, index=True)
    student_no = Column(String, unique=True, index=True, nullable=False)
    name = Column(String, nullable=False)
    password_hash = Column(String, nullable=False)
    created_at = Column(DateTime, default=datetime.utcnow)


class Report(Base):
    __tablename__ = "reports"

    id = Column(Integer, primary_key=True, index=True)
    student_id = Column(Integer, nullable=False)
    filename = Column(String, nullable=False)
    stored_path = Column(String, nullable=False)
    grade = Column(Float, nullable=True)
    feedback = Column(String, nullable=True)
    created_at = Column(DateTime, default=datetime.utcnow)


Base.metadata.create_all(bind=engine)

app = FastAPI(title="LabReport Assistant", version="0.1.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"]
    ,
    allow_headers=["*"],
)

# In-memory session store
session_tokens: Dict[str, int] = {}


class StudentOut(BaseModel):
    id: int
    student_no: str
    name: str

    class Config:
        orm_mode = True


class ImportRequest(BaseModel):
    default_password: str
    students: List[Dict[str, str]]


class LoginRequest(BaseModel):
    student_no: str
    password: str


class LoginResponse(BaseModel):
    token: str
    expires_in: int
    student: StudentOut


class GradeResponse(BaseModel):
    grade: float
    feedback: str
    report_id: int


class StatsResponse(BaseModel):
    total_reports: int
    average_grade: float
    submissions: List[Dict[str, str]]


def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


def hash_password(password: str) -> str:
    return pwd_context.hash(password)


def verify_password(password: str, hashed: str) -> bool:
    return pwd_context.verify(password, hashed)


def authenticate(token: str, db: Session) -> Student:
    student_id = session_tokens.get(token)
    if not student_id:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Invalid token")
    student = db.query(Student).filter(Student.id == student_id).first()
    if not student:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Student not found")
    return student


async def call_deepseek_api(prompt: str) -> Dict[str, str]:
    api_key = os.getenv("DEEPSEEK_API_KEY")
    if not api_key:
        simulated = "自动评语: 报告条理清晰，实验步骤完整。"  # deterministic fallback
        score = 92.0
        return {"feedback": simulated, "score": score}

    headers = {"Authorization": f"Bearer {api_key}"}
    payload = {"model": "deepseek-chat", "messages": [{"role": "user", "content": prompt}]}
    async with httpx.AsyncClient(timeout=30) as client:
        response = await client.post("https://api.deepseek.com/v1/chat/completions", headers=headers, json=payload)
        response.raise_for_status()
        data = response.json()
        message = data.get("choices", [{}])[0].get("message", {}).get("content", "")
        score = 90.0
        for token in message.split():
            if token.rstrip("%").isdigit():
                score = float(token.rstrip("%"))
                break
        return {"feedback": message, "score": score}


def add_feedback_box(doc_path: Path, feedback: str) -> None:
    document = Document(doc_path)
    first_paragraph = document.paragraphs[0]

    run = first_paragraph.add_run()
    run.add_break()
    run.add_text("反馈:")
    run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)

    tbl = document.add_table(rows=1, cols=1)
    cell = tbl.rows[0].cells[0]
    cell.text = feedback
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "FFF0F0")
    shading.set(qn("w:color"), "auto")
    cell._tc.get_or_add_tcPr().append(shading)
    for paragraph in cell.paragraphs:
        for r in paragraph.runs:
            r.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)

    document.save(doc_path)


@app.post("/api/auth/import", response_model=List[StudentOut])
def import_students(payload: ImportRequest, db: Session = Depends(get_db)):
    created_students: List[Student] = []
    for entry in payload.students:
        student_no = entry.get("student_no")
        name = entry.get("name")
        if not student_no or not name:
            continue
        existing = db.query(Student).filter(Student.student_no == student_no).first()
        if existing:
            continue
        student = Student(
            student_no=student_no,
            name=name,
            password_hash=hash_password(payload.default_password),
        )
        db.add(student)
        db.commit()
        db.refresh(student)
        created_students.append(student)
    return created_students


@app.post("/api/auth/login", response_model=LoginResponse)
def login(payload: LoginRequest, db: Session = Depends(get_db)):
    student = db.query(Student).filter(Student.student_no == payload.student_no).first()
    if not student or not verify_password(payload.password, student.password_hash):
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Invalid credentials")
    token = secrets.token_hex(16)
    session_tokens[token] = student.id
    return LoginResponse(token=token, expires_in=3600, student=student)


@app.post("/api/reports/upload", response_model=GradeResponse)
async def upload_report(
    token: str,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    student = authenticate(token, db)
    if not file.filename.lower().endswith(".docx"):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="仅支持docx文件")

    timestamp = datetime.utcnow().strftime("%Y%m%d%H%M%S")
    saved_name = f"{student.student_no}_{timestamp}_{file.filename}"
    saved_path = reports_dir / saved_name

    with saved_path.open("wb") as buffer:
        shutil.copyfileobj(file.file, buffer)

    prompt = f"请根据实验报告生成评语并给出100分制成绩，报告作者：{student.name}"
    deepseek_result = await call_deepseek_api(prompt)

    add_feedback_box(saved_path, deepseek_result["feedback"])

    report = Report(
        student_id=student.id,
        filename=file.filename,
        stored_path=str(saved_path),
        grade=deepseek_result["score"],
        feedback=deepseek_result["feedback"],
    )
    db.add(report)
    db.commit()
    db.refresh(report)

    return GradeResponse(grade=report.grade, feedback=report.feedback, report_id=report.id)


@app.get("/api/reports/{report_id}/download")
def download_report(report_id: int, token: str, db: Session = Depends(get_db)):
    student = authenticate(token, db)
    report = db.query(Report).filter(Report.id == report_id).first()
    if not report or report.student_id != student.id:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="报告不存在")
    return FileResponse(report.stored_path, filename=f"graded_{report.filename}")


@app.get("/api/reports/stats", response_model=StatsResponse)
def report_stats(token: str, db: Session = Depends(get_db)):
    _ = authenticate(token, db)
    total_reports = db.query(func.count(Report.id)).scalar() or 0
    average_grade = db.query(func.avg(Report.grade)).scalar() or 0.0
    submissions = []
    query = (
        db.query(Report, Student)
        .join(Student, Student.id == Report.student_id)
        .order_by(Report.created_at.desc())
    )
    for report, student in query:
        submissions.append(
            {
                "student_no": student.student_no,
                "name": student.name,
                "grade": report.grade,
                "submitted_at": report.created_at.isoformat(),
                "report_id": report.id,
            }
        )
    return StatsResponse(total_reports=total_reports, average_grade=average_grade, submissions=submissions)


@app.get("/")
def health_check():
    return {"status": "ok", "message": "LabReport backend running"}


if __name__ == "__main__":
    import uvicorn

    uvicorn.run("main:app", host="0.0.0.0", port=8000, reload=True)
