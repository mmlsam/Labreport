import io
import os
import ast
from contextlib import nullcontext
from datetime import datetime
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openai import OpenAI

from app import create_app, db
from app.models import Admin, Experiment, ReportSubmission, Student
from app.routes import validated_storage_path


class StoragePathTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name) / "uploads"
        self.root.mkdir()

    def test_normal_path(self):
        self.assertEqual(validated_storage_path(str(self.root), "report.docx"),
                         self.root / "report.docx")

    def test_application_does_not_call_python39_path_api(self):
        tree = ast.parse((Path(__file__).parents[1] / "app" / "routes.py").read_text())
        self.assertFalse(any(isinstance(node, ast.Attribute) and node.attr == "is_relative_to"
                             for node in ast.walk(tree)))

    def test_llm_client_dependency_compatibility(self):
        # This test must not depend on the host's proxy configuration.
        with patch.dict(os.environ, {}, clear=True):
            client = OpenAI(api_key="test-only")
            client.close()

    def test_reject_escape_and_root(self):
        for name in ("../outside.docx", "../uploads-other/file.docx",
                     str(self.root.parent / "absolute.docx"), ".", ""):
            with self.subTest(name=name), self.assertRaises(ValueError):
                validated_storage_path(str(self.root), name)

    def test_normalized_inside_path(self):
        self.assertEqual(validated_storage_path(str(self.root), "sub/../report.docx"),
                         self.root / "report.docx")

    def test_symlink_file_escape(self):
        target = self.root.parent / "outside.docx"
        target.write_bytes(b"unchanged")
        (self.root / "link.docx").symlink_to(target)
        with self.assertRaises(ValueError):
            validated_storage_path(str(self.root), "link.docx")
        self.assertEqual(target.read_bytes(), b"unchanged")

    def test_symlink_directory_escape(self):
        (self.root / "link").symlink_to(self.root.parent, target_is_directory=True)
        with self.assertRaises(ValueError):
            validated_storage_path(str(self.root), "link/report.docx")

    def test_configured_root_symlink(self):
        alias = self.root.parent / "alias"
        alias.symlink_to(self.root, target_is_directory=True)
        self.assertEqual(validated_storage_path(str(alias), "report.docx"),
                         self.root / "report.docx")


class UploadTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.app = create_app({
            "TESTING": True,
            "SECRET_KEY": "test",
            "SQLALCHEMY_DATABASE_URI": "sqlite://",
            "UPLOAD_FOLDER": str(self.root / "uploads"),
            "PROCESSED_FOLDER": str(self.root / "processed"),
            "WTF_CSRF_ENABLED": False,
        })
        self.context = self.app.app_context()
        self.context.push()
        self.addCleanup(self.context.pop)
        self.addCleanup(db.session.remove)
        student = Student(student_number="../../", name="Test", password_hash="unused")
        other = Student(student_number="other", name="Other", password_hash="unused")
        experiment = Experiment(name="Test experiment")
        db.session.add_all([student, other, experiment,
                            Admin(username="test-admin", password_hash="unused")])
        db.session.commit()
        self.student_id, self.other_id, self.experiment_id = student.id, other.id, experiment.id
        self.client = self.app.test_client()
        with self.client.session_transaction() as session:
            session["role"] = "student"
            session["student_id"] = student.id

    def document(self):
        stream = io.BytesIO()
        doc = Document()
        doc.add_paragraph("实验目的 实验原理 实验内容 程序 结果 总结 心得")
        doc.save(stream)
        stream.seek(0)
        return stream

    def upload(self, name="report.docx"):
        return self.client.post("/student/submit/%s" % self.experiment_id,
                                data={"file": (self.document(), name)})

    def test_real_upload_without_python39_path_api(self):
        # Python 3.12's relative_to internally calls is_relative_to. Only install
        # the sentinel on runtimes where the latter does not exist.
        guard = nullcontext() if hasattr(Path, "is_relative_to") else patch.object(
            Path, "is_relative_to", create=True,
            side_effect=AssertionError("Python 3.9 API used"))
        with guard:
            response = self.upload("../../报告.docx")
        self.assertEqual(response.status_code, 302)
        self.assertTrue(response.location.endswith("/student/dashboard"))
        submission = ReportSubmission.query.one()
        stored, processed = Path(submission.stored_path), Path(submission.processed_path)
        self.assertEqual(stored.parent, self.root / "uploads")
        self.assertEqual(processed.parent, self.root / "processed")
        self.assertTrue(stored.name.startswith("%s_" % self.student_id))
        self.assertTrue(stored.is_file())
        self.assertIn("成绩", Document(processed).paragraphs[0].text)

    def test_same_second_resubmission_does_not_overwrite_files(self):
        with patch("app.routes.datetime") as clock:
            clock.utcnow.return_value = datetime(2026, 1, 1)
            self.upload()
            original = ReportSubmission.query.one().stored_path
            self.upload()
        self.assertEqual(ReportSubmission.query.count(), 1)
        self.assertNotEqual(original, ReportSubmission.query.one().stored_path)
        self.assertEqual(len(list((self.root / "uploads").iterdir())), 2)

    def test_reject_invalid_extension(self):
        self.assertEqual(self.upload("report.exe").status_code, 302)
        self.assertEqual(ReportSubmission.query.count(), 0)
        self.assertEqual(list((self.root / "uploads").iterdir()), [])

    def test_missing_file(self):
        response = self.client.post("/student/submit/%s" % self.experiment_id)
        self.assertEqual(response.status_code, 302)
        self.assertEqual(ReportSubmission.query.count(), 0)

    def test_authentication_required(self):
        with self.client.session_transaction() as session:
            session.clear()
        self.assertTrue(self.upload().location.endswith("/student/login"))
        self.assertEqual(ReportSubmission.query.count(), 0)

    def test_csrf_required(self):
        self.app.config["WTF_CSRF_ENABLED"] = True
        self.assertEqual(self.upload().status_code, 400)
        self.assertEqual(ReportSubmission.query.count(), 0)

    def test_size_limit(self):
        self.app.config["MAX_CONTENT_LENGTH"] = 100
        self.assertEqual(self.upload().status_code, 413)
        self.assertEqual(ReportSubmission.query.count(), 0)

    def test_missing_experiment(self):
        response = self.client.post("/student/submit/99999",
                                    data={"file": (self.document(), "report.docx")})
        self.assertEqual(response.status_code, 404)
        self.assertEqual(ReportSubmission.query.count(), 0)

    def test_both_destinations_checked_before_save(self):
        for folder in ("UPLOAD_FOLDER", "PROCESSED_FOLDER"):
            with self.subTest(folder=folder), patch("app.routes.datetime") as clock, \
                    patch("app.routes.uuid4") as token, \
                    patch("app.routes.evaluate_report") as evaluate:
                clock.utcnow.return_value.strftime.return_value = "fixed"
                token.return_value.hex = "token"
                name = "%s_%s_fixed_token_report.docx" % (self.student_id, self.experiment_id)
                outside = self.root / "outside.docx"
                outside.write_bytes(b"unchanged")
                link = Path(self.app.config[folder]) / name
                link.symlink_to(outside)
                response = self.upload()
                self.assertEqual(response.status_code, 302)
                self.assertTrue(response.location.endswith("/student/submit/%s" % self.experiment_id))
                self.assertEqual(outside.read_bytes(), b"unchanged")
                self.assertEqual(ReportSubmission.query.count(), 0)
                evaluate.assert_not_called()
                link.unlink()
                self.assertEqual(list((self.root / "uploads").iterdir()), [])

    def test_download_ownership(self):
        self.upload()
        submission = ReportSubmission.query.one()
        response = self.client.get("/student/download/%s" % submission.id)
        self.assertEqual(response.status_code, 200)
        response.close()
        with self.client.session_transaction() as session:
            session["student_id"] = self.other_id
        response = self.client.get("/student/download/%s" % submission.id)
        self.assertEqual(response.status_code, 302)
        self.assertNotIn("Content-Disposition", response.headers)
        self.assertTrue(response.location.endswith("/student/dashboard"))


if __name__ == "__main__":
    unittest.main()
